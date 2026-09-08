"""Binary asset storage and profile upload validation tests."""

import asyncio
from io import BytesIO
from pathlib import Path

import docx
import pptx
import pytest
from PIL import Image
from pypdf import PdfWriter

from launchkit.assets.storage import LocalAssetBlobStore
from launchkit.assets.validation import (
    UploadTooLargeError,
    UploadValidationError,
    safe_filename,
    validate_brand_upload,
    validate_profile_upload,
)


def document_bytes(kind: str) -> bytes:
    buffer = BytesIO()
    if kind == "docx":
        document = docx.Document()
        document.add_paragraph("Company")
        document.save(buffer)
    elif kind == "pptx":
        presentation = pptx.Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[0])
        presentation.save(buffer)
    elif kind == "pdf":
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.write(buffer)
    elif kind in {"png", "jpg"}:
        Image.new("RGB", (2, 2), "white").save(buffer, format="PNG" if kind == "png" else "JPEG")
    return buffer.getvalue()


@pytest.mark.parametrize(
    ("filename", "content_type", "content"),
    [
        ("profile.pdf", "application/pdf", document_bytes("pdf")),
        (
            "profile.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            document_bytes("docx"),
        ),
        (
            "profile.pptx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            document_bytes("pptx"),
        ),
        ("profile.txt", "text/plain", b"Company"),
        ("profile.md", "text/markdown", b"# Company"),
        ("profile.png", "image/png", document_bytes("png")),
        ("profile.jpg", "image/jpeg", document_bytes("jpg")),
    ],
    ids=["pdf", "docx", "pptx", "txt", "markdown", "png", "jpeg"],
)
def test_profile_upload_accepts_advertised_types(
    filename: str, content_type: str, content: bytes
) -> None:
    validated = validate_profile_upload(filename, content_type, content, max_bytes=1024 * 1024)

    assert validated.filename == filename
    assert validated.content == content


def test_logo_upload_accepts_svg() -> None:
    validated = validate_brand_upload(
        "logo.svg",
        "image/svg+xml",
        b'<svg xmlns="http://www.w3.org/2000/svg"><circle r="1"/></svg>',
        kind="logo",
    )
    assert validated.extension == "svg"


def test_logo_upload_rejects_script_svg() -> None:
    with pytest.raises(UploadValidationError):
        validate_brand_upload(
            "logo.svg",
            "image/svg+xml",
            b"<svg><script>alert(1)</script></svg>",
            kind="logo",
        )


@pytest.mark.parametrize(
    ("filename", "content_type", "content", "error"),
    [
        ("profile.csv", "text/csv", b"x", UploadValidationError),
        ("profile.pdf", "text/plain", b"x", UploadValidationError),
        ("profile.txt", "text/plain", b"", UploadValidationError),
        ("profile.png", "image/png", b"not-an-image", UploadValidationError),
        (
            "profile.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            b"bad",
            UploadValidationError,
        ),
        ("profile.txt", "text/plain", b"too large", UploadTooLargeError),
    ],
    ids=["extension", "mime", "empty", "image", "office", "size"],
)
def test_profile_upload_rejects_invalid_content(
    filename: str,
    content_type: str,
    content: bytes,
    error: type[Exception],
) -> None:
    max_bytes = 2 if error is UploadTooLargeError else 1024
    with pytest.raises(error):
        validate_profile_upload(filename, content_type, content, max_bytes=max_bytes)


def test_project_asset_key_scopes_under_owner() -> None:
    from launchkit.assets.storage import project_asset_key

    assert project_asset_key("user/abc!", "prj_1", "logos", "a.png") == (
        "user-abc/projects/prj_1/logos/a.png"
    )
    assert project_asset_key("", "prj_1", "x") == "owner/projects/prj_1/x"


def test_local_asset_store_round_trip_and_traversal_protection(tmp_path: Path) -> None:
    store = LocalAssetBlobStore(tmp_path)

    async def scenario() -> bytes:
        await store.put("projects/one/file.txt", b"content", "text/plain")
        return await store.get("projects/one/file.txt")

    assert asyncio.run(scenario()) == b"content"
    with pytest.raises(ValueError, match="escapes"):
        asyncio.run(store.get("../outside.txt"))


def test_local_asset_store_missing_key_explains_shared_storage(tmp_path: Path) -> None:
    store = LocalAssetBlobStore(tmp_path)

    with pytest.raises(FileNotFoundError, match="share local_data"):
        asyncio.run(store.get("projects/missing/file.pdf"))


def test_filename_normalization_removes_paths_headers_and_unicode() -> None:
    assert safe_filename('../folder/brand"\n.png') == "brand_.png"
    assert safe_filename("C:\\folder\\logo.png") == "logo.png"
    assert safe_filename("profile-شركة.pdf") == "profile-_.pdf"
